"""
Génère un fichier Excel d'exemple pour tester l'import dans le module
de calcul des commissions ABC DIS.

3 feuilles :
  - Objectifs       (periode, carte, matricule, target)
  - Realisations    (periode, carte, matricule, ca_realise)
  - NotesTriage     (periode, matricule, note)

Sortie : C:/Projets/RH-Commission/frontend/public/samples/ventes_mai_2026.xlsx
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

PERIODE = "2026-05"
CARTES = ["Coca Cola", "Wall's", "Ferrero Rocher"]
MATRICULES = ["P001", "P002", "P003", "P004", "P005",
              "P006", "P007", "P008", "P009"]

OBJECTIFS = [
    (PERIODE, "Coca Cola",      "P001", 50000),
    (PERIODE, "Coca Cola",      "P002", 45000),
    (PERIODE, "Coca Cola",      "P003", 60000),
    (PERIODE, "Wall's",          "P004", 30000),
    (PERIODE, "Wall's",          "P005", 25000),
    (PERIODE, "Wall's",          "P006", 28000),
    (PERIODE, "Ferrero Rocher", "P007", 40000),
    (PERIODE, "Ferrero Rocher", "P008", 35000),
    (PERIODE, "Ferrero Rocher", "P009", 42000),
]

REALISATIONS = [
    (PERIODE, "Coca Cola",      "P001", 52000),
    (PERIODE, "Coca Cola",      "P002", 41000),
    (PERIODE, "Coca Cola",      "P003", 68000),
    (PERIODE, "Wall's",          "P004", 31500),
    (PERIODE, "Wall's",          "P005", 22000),
    (PERIODE, "Wall's",          "P006", 30200),
    (PERIODE, "Ferrero Rocher", "P007", 38000),
    (PERIODE, "Ferrero Rocher", "P008", 36500),
    (PERIODE, "Ferrero Rocher", "P009", 44000),
]

NOTES_TRIAGE = [
    (PERIODE, "P001", 88),
    (PERIODE, "P002", 72),
    (PERIODE, "P003", 95),
    (PERIODE, "P004", 81),
    (PERIODE, "P005", 65),
    (PERIODE, "P006", 79),
    (PERIODE, "P007", 90),
    (PERIODE, "P008", 84),
    (PERIODE, "P009", 92),
]


def style_header(ws, n_cols):
    fill = PatternFill("solid", fgColor="111827")
    font = Font(bold=True, color="FFFFFF", size=11)
    align = Alignment(horizontal="center", vertical="center")
    thin = Side(border_style="thin", color="E5E7EB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for col in range(1, n_cols + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = fill
        cell.font = font
        cell.alignment = align
        cell.border = border
    ws.row_dimensions[1].height = 28


def autosize(ws):
    for col_cells in ws.columns:
        length = max(len(str(c.value or "")) for c in col_cells)
        ws.column_dimensions[col_cells[0].column_letter].width = max(length + 4, 14)


wb = Workbook()

# === Sheet 1 — Objectifs ===
ws = wb.active
ws.title = "Objectifs"
ws.append(["periode", "carte", "matricule", "target"])
for row in OBJECTIFS:
    ws.append(list(row))
style_header(ws, 4)
autosize(ws)

# === Sheet 2 — Realisations ===
ws = wb.create_sheet("Realisations")
ws.append(["periode", "carte", "matricule", "ca_realise"])
for row in REALISATIONS:
    ws.append(list(row))
style_header(ws, 4)
autosize(ws)

# === Sheet 3 — NotesTriage ===
ws = wb.create_sheet("NotesTriage")
ws.append(["periode", "matricule", "note"])
for row in NOTES_TRIAGE:
    ws.append(list(row))
style_header(ws, 3)
autosize(ws)

OUT = r"C:\Projets\RH-Commission\frontend\public\samples\ventes_mai_2026.xlsx"
wb.save(OUT)
print(f"OK : {OUT}")
